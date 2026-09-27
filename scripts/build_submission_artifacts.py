from pathlib import Path
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch, mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.pdfgen import canvas
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
DOCS.mkdir(parents=True, exist_ok=True)

NAVY = RGBColor(8, 18, 38)
NAVY2 = RGBColor(12, 26, 52)
BLUE = RGBColor(49, 143, 255)
CYAN = RGBColor(74, 220, 229)
GOLD = RGBColor(246, 187, 67)
WHITE = RGBColor(245, 248, 252)
MUTED = RGBColor(173, 190, 213)
GREEN = RGBColor(77, 207, 154)

TITLE = "SnapLearn Edge"
TAGLINE = "Privacy-first on-device AI lecture & meeting copilot for Snapdragon-powered HP PCs"

def set_bg(slide, color=NAVY):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color

def add_text(slide, text, x, y, w, h, size=22, color=WHITE, bold=False, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.name = "Aptos"
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return tb

def rounded(slide, x, y, w, h, fill_rgb, line_rgb=None, radius=True):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_rgb
    shape.line.color.rgb = line_rgb or fill_rgb
    return shape

def card(slide, x, y, w, h, eyebrow, title, body, accent=BLUE):
    rounded(slide, x, y, w, h, NAVY2, RGBColor(30, 51, 82))
    rounded(slide, x, y, 0.08, h, accent, accent, radius=False)
    add_text(slide, eyebrow.upper(), x + 0.22, y + 0.12, w - 0.35, 0.3, 9, accent, True)
    add_text(slide, title, x + 0.22, y + 0.48, w - 0.35, 0.52, 18, WHITE, True)
    add_text(slide, body, x + 0.22, y + 1.0, w - 0.35, h - 1.12, 11.5, MUTED)

def make_pptx(path: Path):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    s = prs.slides.add_slide(blank)
    set_bg(s)
    for i, (x, y, w) in enumerate([(8.8, 1.0, 3.5), (9.35, 1.75, 2.7), (8.25, 2.5, 4.0), (9.1, 3.25, 3.3), (8.55, 4.0, 3.8)]):
        line = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.025))
        line.fill.solid()
        line.fill.fore_color.rgb = CYAN if i % 2 == 0 else GOLD
        line.line.fill.background()
    rounded(s, 9.6, 1.4, 2.5, 2.5, RGBColor(15, 38, 72), CYAN)
    add_text(s, "AI", 9.6, 1.8, 2.5, 1.1, 44, WHITE, True, align=PP_ALIGN.CENTER)
    add_text(s, "ON DEVICE", 9.6, 2.75, 2.5, 0.4, 10, CYAN, True, align=PP_ALIGN.CENTER)
    add_text(s, "SNAPDRAGON AI LAB • BUILD & PRESENT CHALLENGE", 0.75, 0.65, 7.3, 0.35, 10, CYAN, True)
    add_text(s, TITLE, 0.75, 1.52, 7.5, 0.9, 42, WHITE, True)
    add_text(s, TAGLINE, 0.75, 2.46, 6.9, 1.0, 19, MUTED)
    rounded(s, 0.75, 4.15, 2.05, 0.56, BLUE, BLUE)
    add_text(s, "OFFLINE-FIRST", 0.75, 4.15, 2.05, 0.56, 10, WHITE, True, align=PP_ALIGN.CENTER)
    rounded(s, 2.95, 4.15, 1.95, 0.56, RGBColor(29, 91, 86), RGBColor(29, 91, 86))
    add_text(s, "PRIVACY-FIRST", 2.95, 4.15, 1.95, 0.56, 10, WHITE, True, align=PP_ALIGN.CENTER)
    rounded(s, 5.05, 4.15, 2.55, 0.56, RGBColor(78, 59, 26), RGBColor(78, 59, 26))
    add_text(s, "SNAPDRAGON-READY", 5.05, 4.15, 2.55, 0.56, 10, WHITE, True, align=PP_ALIGN.CENTER)
    add_text(s, "Pratham Priyanshu Mohanty", 0.75, 6.36, 5, 0.35, 12, WHITE, True)
    add_text(s, "Solution Submission Round", 0.75, 6.72, 5, 0.3, 10, MUTED)

    s = prs.slides.add_slide(blank)
    set_bg(s)
    add_text(s, "01  PROBLEM → SOLUTION → USER VALUE", 0.65, 0.36, 6.2, 0.35, 10, CYAN, True)
    add_text(s, "Turn every lecture or meeting into private, actionable knowledge.", 0.65, 0.76, 11.9, 0.65, 26, WHITE, True)
    card(s, 0.65, 1.65, 3.85, 2.12, "Problem", "Cloud-first AI creates friction",
         "Sensitive recordings may leave the PC. Connectivity, latency and recurring inference cost can interrupt learning and work.", BLUE)
    card(s, 4.74, 1.65, 3.85, 2.12, "Solution", "SnapLearn Edge keeps the workflow local",
         "Audio → transcript → smart notes → action items → revision questions → grounded Q&A, designed for Snapdragon-powered HP PCs.", CYAN)
    card(s, 8.83, 1.65, 3.85, 2.12, "Value", "One private productivity loop",
         "Students and knowledge workers get usable outputs even when connectivity is weak—while preserving an offline-first architecture.", GOLD)
    add_text(s, "USER FLOW", 0.65, 4.15, 1.2, 0.3, 10, MUTED, True)
    labels = [
        ("1", "CAPTURE", "Lecture / meeting audio"),
        ("2", "TRANSCRIBE", "On-device Whisper ASR"),
        ("3", "UNDERSTAND", "Local LLM structures notes"),
        ("4", "ACT", "Ask, revise, export")
    ]
    x = 0.65
    for i, (n, t, b) in enumerate(labels):
        rounded(s, x, 4.62, 2.88, 1.5, NAVY2, RGBColor(35, 60, 92))
        accent = BLUE if i < 2 else CYAN
        rounded(s, x + 0.18, 4.82, 0.45, 0.45, accent, accent)
        add_text(s, n, x + 0.18, 4.82, 0.45, 0.45, 11, WHITE, True, align=PP_ALIGN.CENTER)
        add_text(s, t, x + 0.76, 4.74, 1.85, 0.35, 13, WHITE, True)
        add_text(s, b, x + 0.76, 5.16, 1.85, 0.62, 10, MUTED)
        x += 3.05
    add_text(s, "Innovation: not another cloud note-taker—an edge-AI knowledge companion designed around local inference.", 0.65, 6.55, 12.0, 0.34, 11, GREEN, True)

    s = prs.slides.add_slide(blank)
    set_bg(s)
    add_text(s, "02  TECHNICAL IMPLEMENTATION & CHALLENGE FIT", 0.65, 0.36, 7.0, 0.35, 10, CYAN, True)
    add_text(s, "A modular Snapdragon deployment path with honest prototype boundaries.", 0.65, 0.76, 11.7, 0.65, 26, WHITE, True)
    add_text(s, "TARGET AI PIPELINE", 0.65, 1.65, 2.2, 0.3, 10, MUTED, True)
    stages = [
        ("AUDIO", "Upload / capture", BLUE),
        ("ASR", "Whisper-Small-Quantized", CYAN),
        ("CONTEXT", "Chunk + structure", GREEN),
        ("LLM", "Llama 3.2 3B Instruct", GOLD),
        ("OUTPUT", "Notes • Q&A • actions", BLUE)
    ]
    x = 0.65
    for i, (a, b, c) in enumerate(stages):
        rounded(s, x, 2.08, 2.28, 1.27, NAVY2, RGBColor(35, 60, 92))
        add_text(s, a, x + 0.18, 2.26, 1.9, 0.28, 10, c, True)
        add_text(s, b, x + 0.18, 2.60, 1.9, 0.48, 11.5, WHITE, True)
        if i < 4:
            add_text(s, "→", x + 2.31, 2.45, 0.35, 0.4, 18, MUTED, True, align=PP_ALIGN.CENTER)
        x += 2.48
    card(s, 0.65, 3.75, 2.87, 2.1, "Technical", "NPU-oriented design",
         "Model adapters are separated from the Streamlit UI. Qualcomm AI Hub models can be exported/compiled for the target runtime and device.", BLUE)
    card(s, 3.72, 3.75, 2.87, 2.1, "Innovation", "Private edge workflow",
         "Local speech + local generation reduces cloud dependency and keeps sensitive classroom or meeting context on the PC.", CYAN)
    card(s, 6.79, 3.75, 2.87, 2.1, "Deployment", "Evaluator-friendly",
         "Simple Windows/Python setup, model folders kept external, and a documented Snapdragon integration + benchmarking path.", GREEN)
    card(s, 9.86, 3.75, 2.82, 2.1, "Documentation", "Submission-ready",
         "README, architecture, model plan, deployment guide, challenge mapping, project PDF and pitch assets are included.", GOLD)
    add_text(s, "Current status: functional UI + adapter scaffold. Hardware-specific NPU inference and performance numbers must be validated on the target HP Snapdragon PC before claiming benchmarks.", 0.65, 6.25, 12.05, 0.6, 10.5, MUTED)
    prs.save(path)

def make_pitch_pdf(path: Path):
    W, H = landscape((13.333 * inch, 7.5 * inch))
    c = canvas.Canvas(str(path), pagesize=(W, H))
    def bg():
        c.setFillColor(colors.HexColor("#081226"))
        c.rect(0, 0, W, H, fill=1, stroke=0)
    def txt(t, x, y, size=18, color="#F5F8FC", bold=False):
        c.setFillColor(colors.HexColor(color))
        c.setFont("Helvetica-Bold" if bold else "Helvetica", size)
        c.drawString(x, y, t)
    def box(x, y, w, h, head, title, body, accent):
        c.setFillColor(colors.HexColor("#0C1A34"))
        c.roundRect(x, y, w, h, 10, fill=1, stroke=0)
        c.setFillColor(colors.HexColor(accent))
        c.rect(x, y, 5, h, fill=1, stroke=0)
        txt(head.upper(), x + 16, y + h - 24, 8, accent, True)
        txt(title, x + 16, y + h - 50, 13, "#F5F8FC", True)
        st = ParagraphStyle("card", fontName="Helvetica", fontSize=8.5, leading=11, textColor=colors.HexColor("#ADBED5"))
        p = Paragraph(body, st)
        p.wrapOn(c, w - 30, h - 70)
        p.drawOn(c, x + 16, y + 18)

    bg()
    txt("SNAPDRAGON AI LAB • BUILD & PRESENT CHALLENGE", 54, H - 58, 9, "#4ADCE5", True)
    txt(TITLE, 54, H - 132, 34, "#F5F8FC", True)
    p = Paragraph(TAGLINE, ParagraphStyle("tag", fontName="Helvetica", fontSize=16, leading=21, textColor=colors.HexColor("#ADBED5")))
    p.wrapOn(c, 520, 100); p.drawOn(c, 54, H - 226)
    txt("OFFLINE-FIRST   •   PRIVACY-FIRST   •   SNAPDRAGON-READY", 54, H - 312, 10, "#4ADCE5", True)
    txt("Pratham Priyanshu Mohanty", 54, 62, 11, "#F5F8FC", True)
    txt("Solution Submission Round", 54, 44, 9, "#ADBED5")
    c.showPage()

    bg()
    txt("01  PROBLEM → SOLUTION → USER VALUE", 47, H - 45, 9, "#4ADCE5", True)
    txt("Turn every lecture or meeting into private, actionable knowledge.", 47, H - 86, 21, "#F5F8FC", True)
    box(47, H - 275, 260, 132, "Problem", "Cloud-first AI creates friction",
        "Sensitive recordings may leave the PC. Connectivity, latency and recurring inference cost can interrupt learning and work.", "#318FFF")
    box(327, H - 275, 260, 132, "Solution", "Keep the workflow local",
        "Audio → transcript → smart notes → action items → revision questions → grounded Q&A, designed for Snapdragon-powered HP PCs.", "#4ADCE5")
    box(607, H - 275, 260, 132, "Value", "One private productivity loop",
        "Useful outputs even with weak connectivity, while maintaining an offline-first architecture.", "#F6BB43")
    txt("USER FLOW", 47, H - 316, 9, "#ADBED5", True)
    flow = [("1", "CAPTURE", "Lecture / meeting audio"), ("2", "TRANSCRIBE", "On-device Whisper ASR"), ("3", "UNDERSTAND", "Local LLM structures notes"), ("4", "ACT", "Ask, revise, export")]
    x = 47
    for n, t, b in flow:
        c.setFillColor(colors.HexColor("#0C1A34")); c.roundRect(x, H - 435, 192, 82, 8, fill=1, stroke=0)
        txt(n, x + 12, H - 382, 10, "#4ADCE5", True); txt(t, x + 36, H - 382, 10, "#F5F8FC", True); txt(b, x + 36, H - 405, 8, "#ADBED5")
        x += 205
    txt("Innovation: an edge-AI knowledge companion designed around local inference.", 47, 48, 10, "#4DCF9A", True)
    c.showPage()

    bg()
    txt("02  TECHNICAL IMPLEMENTATION & CHALLENGE FIT", 47, H - 45, 9, "#4ADCE5", True)
    txt("A modular Snapdragon deployment path with honest prototype boundaries.", 47, H - 86, 21, "#F5F8FC", True)
    stages = [("AUDIO", "Upload / capture"), ("ASR", "Whisper-Small-Quantized"), ("CONTEXT", "Chunk + structure"), ("LLM", "Llama 3.2 3B Instruct"), ("OUTPUT", "Notes • Q&A • actions")]
    x = 47
    for a, b in stages:
        c.setFillColor(colors.HexColor("#0C1A34")); c.roundRect(x, H - 220, 153, 74, 8, fill=1, stroke=0)
        txt(a, x + 10, H - 173, 8, "#4ADCE5", True); txt(b, x + 10, H - 197, 8.5, "#F5F8FC", True); x += 164
    box(47, H - 430, 196, 146, "Technical", "NPU-oriented design", "Model adapters are separated from the UI and documented for Qualcomm AI Hub export/compile workflows.", "#318FFF")
    box(255, H - 430, 196, 146, "Innovation", "Private edge workflow", "Local speech + generation reduce cloud dependence and keep sensitive context on the PC.", "#4ADCE5")
    box(463, H - 430, 196, 146, "Deployment", "Evaluator-friendly", "Windows/Python setup with external model folders and a documented device-validation path.", "#4DCF9A")
    box(671, H - 430, 196, 146, "Documentation", "Submission-ready", "README, architecture, model plan, deployment guide, challenge mapping and pitch assets.", "#F6BB43")
    p = Paragraph("Current status: functional UI + adapter scaffold. Hardware-specific NPU inference and performance metrics require validation on the target Snapdragon-powered HP PC.",
                  ParagraphStyle("foot", fontName="Helvetica", fontSize=8.5, leading=11, textColor=colors.HexColor("#ADBED5")))
    p.wrapOn(c, 820, 60); p.drawOn(c, 47, 45)
    c.save()

def make_description_pdf(path: Path):
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="Hero", fontName="Helvetica-Bold", fontSize=24, leading=28, textColor=colors.HexColor("#081226"), spaceAfter=8))
    styles.add(ParagraphStyle(name="Sub", fontName="Helvetica", fontSize=10.5, leading=15, textColor=colors.HexColor("#39516D"), spaceAfter=12))
    styles.add(ParagraphStyle(name="H2x", fontName="Helvetica-Bold", fontSize=13, leading=16, textColor=colors.HexColor("#0B5D9A"), spaceBefore=7, spaceAfter=5))
    styles.add(ParagraphStyle(name="Bodyx", fontName="Helvetica", fontSize=9.7, leading=14, textColor=colors.HexColor("#233449"), spaceAfter=7))
    styles.add(ParagraphStyle(name="Smallx", fontName="Helvetica", fontSize=8.7, leading=12, textColor=colors.HexColor("#52677E"), spaceAfter=5))
    doc = SimpleDocTemplate(str(path), pagesize=A4, rightMargin=16*mm, leftMargin=16*mm, topMargin=15*mm, bottomMargin=15*mm, title="SnapLearn Edge - Project Description")
    story = [Paragraph("SnapLearn Edge", styles["Hero"]), Paragraph(TAGLINE, styles["Sub"])]
    story.append(Paragraph(
        "SnapLearn Edge is a proposed privacy-first AI application for students and knowledge workers. It turns lecture, meeting and voice-note audio into a transcript, structured notes, action items, revision questions and contextual Q&A. The product is designed around an offline-first architecture for Snapdragon-powered HP PCs so that sensitive audio and derived knowledge can remain on the device whenever possible.",
        styles["Bodyx"]
    ))
    data = [
        ["Challenge Area", "How SnapLearn Edge addresses it"],
        ["Technical Implementation", "Modular audio → ASR → transcript → local LLM → notes/Q&A pipeline with separate model-adapter layers."],
        ["Application Use Case & Innovation", "A local AI knowledge companion for lectures and meetings, prioritising privacy, connectivity independence and lower cloud dependence."],
        ["Deployment & Accessibility", "Simple Streamlit UI, Windows/Python setup, external model folders and a documented Snapdragon integration path."],
        ["Presentation & Documentation", "README, architecture, model plan, deployment guide, challenge mapping, project description and pitch deck."]
    ]
    t = Table(data, colWidths=[43*mm, 130*mm])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#081226")), ("TEXTCOLOR", (0,0), (-1,0), colors.white),
        ("FONTNAME", (0,0), (-1,0), "Helvetica-Bold"), ("FONTNAME", (0,1), (0,-1), "Helvetica-Bold"),
        ("FONTSIZE", (0,0), (-1,-1), 8.2), ("LEADING", (0,0), (-1,-1), 11), ("VALIGN", (0,0), (-1,-1), "TOP"),
        ("GRID", (0,0), (-1,-1), 0.4, colors.HexColor("#C9D5E3")), ("BACKGROUND", (0,1), (-1,-1), colors.HexColor("#F6F9FC")),
        ("LEFTPADDING", (0,0), (-1,-1), 6), ("RIGHTPADDING", (0,0), (-1,-1), 6), ("TOPPADDING", (0,0), (-1,-1), 6), ("BOTTOMPADDING", (0,0), (-1,-1), 6)
    ]))
    story += [Spacer(1, 3*mm), t]
    sections = [
        ("Problem", "Cloud-first transcription and generative-AI workflows can require users to upload sensitive recordings, depend on reliable connectivity and pay recurring inference costs. In classrooms and business meetings, those constraints can reduce trust and practical usefulness."),
        ("Proposed solution", "The application provides one private workflow: capture or upload audio, transcribe it locally, structure the transcript into concise notes, identify action items, generate revision questions and answer questions grounded in the session context."),
        ("AI models and Snapdragon optimisation", "The proposed speech layer uses a Qualcomm AI Hub Whisper model such as Whisper-Small-Quantized or Whisper-Base. The local reasoning layer is designed around Llama 3.2 3B Instruct. On Snapdragon X-class Windows PCs, the implementation path targets Qualcomm-compatible acceleration/runtime options and quantised models to reduce memory and power cost."),
        ("User experience", "A simple desktop-friendly interface lets users upload audio, generate notes and ask questions without exposing model complexity. Outputs are intended to be exportable for revision, documentation and follow-up."),
        ("Deployment approach", "The repository keeps third-party model binaries outside source control. Evaluators can install the Python dependencies, place/export compatible model assets into the documented model folders, connect the runtime adapters and launch the Streamlit interface."),
        ("Current prototype status", "The repository includes a functional application UI and model-adapter scaffold. It does not claim measured Snapdragon latency, throughput or power figures. Those values should be added only after hardware-specific NPU inference is integrated and tested on the target HP Snapdragon PC."),
        ("Why it is a strong challenge fit", "SnapLearn Edge directly combines an everyday productivity problem with an edge-AI architecture. The concept is practical, easy to demonstrate, privacy-oriented and intentionally structured around Snapdragon-powered HP PCs and Qualcomm AI Hub/open-source models.")
    ]
    for h, b in sections:
        story += [Paragraph(h, styles["H2x"]), Paragraph(b, styles["Bodyx"])]
    story += [Spacer(1, 2*mm), Paragraph("<b>Author:</b> Pratham Priyanshu Mohanty", styles["Smallx"]),
              Paragraph("<b>Repository:</b> github.com/Pratham2810/SnapLearn-Edge-AI", styles["Smallx"])]
    doc.build(story)

if __name__ == "__main__":
    make_description_pdf(DOCS / "Project_Description.pdf")
    make_pptx(DOCS / "Pitch_Deck.pptx")
    make_pitch_pdf(DOCS / "Pitch_Deck.pdf")
    print("Generated submission artifacts in", DOCS)
